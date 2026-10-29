import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("STT", 6, "center", None),
    ("Ngày nhận mail", 18, "center", None),
    ("Ngày hóa đơn", 14, "center", None),
    ("Mẫu số", 10, "center", None),
    ("Ký hiệu", 12, "center", None),
    ("Số hóa đơn", 14, "center", "@"),
    ("Biển số xe", 14, "center", None),
    ("Tên hàng hóa", 26, "left", None),
    ("ĐVT", 8, "center", None),
    ("Số lượng (Lít)", 14, "right", "#,##0.000"),
    ("Đơn giá (VNĐ)", 14, "right", "#,##0"),
    ("Thành tiền (VNĐ)", 16, "right", "#,##0"),
    ("Tiền thuế VAT", 14, "right", "#,##0"),
    ("Tổng thanh toán (VNĐ)", 20, "right", "#,##0"),
    ("Bên bán / Cửa hàng", 32, "left", None),
    ("Địa chỉ cửa hàng", 32, "left", None),
    ("Bên mua", 28, "left", None),
    ("Mã tra cứu (Fkey)", 18, "center", None),
    ("File PDF", 12, "center", None),
    ("File XML", 12, "center", None)
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Segoe UI", size=10)
LINK_FONT = Font(name="Segoe UI", size=10, color="0563C1", underline="single")

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

class ExcelManager:
    def __init__(self, excel_path):
        self.excel_path = excel_path
        os.makedirs(os.path.dirname(self.excel_path), exist_ok=True)
        self._ensure_workbook()

    def _ensure_workbook(self):
        """Khởi tạo file Excel nếu chưa tồn tại"""
        if not os.path.exists(self.excel_path):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Danh Sách Hóa Đơn"
            ws.views.sheetView[0].showGridLines = True

            # Ghi tiêu đề cột
            for col_idx, (col_name, col_width, align, _) in enumerate(COLUMNS, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.column_dimensions[get_column_letter(col_idx)].width = col_width

            ws.row_dimensions[1].height = 28
            ws.freeze_panes = "A2"
            wb.save(self.excel_path)
            wb.close()

    def get_existing_invoices(self):
        """Lấy danh sách các số hóa đơn và mã tra cứu đã tồn tại trong file để chống trùng"""
        existing_keys = set()
        if not os.path.exists(self.excel_path):
            return existing_keys

        wb = openpyxl.load_workbook(self.excel_path, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) >= 18:
                ky_hieu = str(row[4] or '').strip()
                so_hd = str(row[5] or '').strip()
                fkey = str(row[17] or '').strip()
                if so_hd:
                    existing_keys.add(f"{ky_hieu}_{so_hd}")
                if fkey:
                    existing_keys.add(fkey)
        wb.close()
        return existing_keys

    def get_invoice_count(self):
        """Lấy số lượng dòng hóa đơn thực tế trong file Excel"""
        if not os.path.exists(self.excel_path):
            return 0
        try:
            wb = openpyxl.load_workbook(self.excel_path, read_only=True)
            ws = wb.active
            count = max(0, ws.max_row - 1)
            wb.close()
            return count
        except Exception:
            return 0

    def append_invoices(self, invoice_list):
        """
        Ghi thêm danh sách hóa đơn vào file Excel, tự động bỏ qua các hóa đơn đã tồn tại.
        Trả về số lượng hóa đơn mới được ghi thêm.
        """
        if not invoice_list:
            return 0

        self._ensure_workbook()
        existing_keys = self.get_existing_invoices()

        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb.active

        current_row = ws.max_row
        stt = current_row - 1 if current_row > 1 else 0
        added_count = 0

        for inv in invoice_list:
            ky_hieu = str(inv.get('ky_hieu', '')).strip()
            so_hd = str(inv.get('so_hd', '')).strip()
            fkey = str(inv.get('ma_tra_cuu', '')).strip()

            key1 = f"{ky_hieu}_{so_hd}"
            if key1 in existing_keys or fkey in existing_keys:
                print(f"[ExcelManager] Hóa đơn {ky_hieu}-{so_hd} (Fkey: {fkey}) đã có trong file, bỏ qua.")
                continue

            current_row += 1
            stt += 1
            added_count += 1

            pdf_path = inv.get('pdf_path', '')
            xml_path = inv.get('xml_path', '')

            # Chuẩn bị dữ liệu từng ô
            row_data = [
                stt,
                inv.get('received_time', ''),
                inv.get('ngay_lap', ''),
                inv.get('mau_so', ''),
                ky_hieu,
                so_hd,
                inv.get('bien_so_xe', ''),
                inv.get('ten_hang', ''),
                inv.get('dvt', 'Lít'),
                inv.get('so_luong', 0.0),
                inv.get('don_gia', 0.0),
                inv.get('thanh_tien', 0.0),
                inv.get('tien_thue', 0.0),
                inv.get('tong_tien_tt', 0.0),
                inv.get('ten_ben_ban', ''),
                inv.get('dia_chi_ddkd', '') or inv.get('dia_chi_ban', ''),
                inv.get('ten_ben_mua', ''),
                fkey,
                "Xem PDF" if pdf_path else "",
                "Xem XML" if xml_path else ""
            ]

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER

                # Căn lề và định dạng số
                _, _, align, num_fmt = COLUMNS[col_idx - 1]
                cell.alignment = Alignment(horizontal=align, vertical="center")
                if num_fmt:
                    cell.number_format = num_fmt

                # Gán Hyperlink cho PDF
                if col_idx == 19 and pdf_path and os.path.exists(pdf_path):
                    cell.hyperlink = pdf_path
                    cell.font = LINK_FONT

                # Gán Hyperlink cho XML
                if col_idx == 20 and xml_path and os.path.exists(xml_path):
                    cell.hyperlink = xml_path
                    cell.font = LINK_FONT

            ws.row_dimensions[current_row].height = 22
            existing_keys.add(key1)
            if fkey:
                existing_keys.add(fkey)

        wb.save(self.excel_path)
        wb.close()
        print(f"[ExcelManager] Đã cập nhật thành công {added_count} hóa đơn mới vào {self.excel_path}")
        return added_count
